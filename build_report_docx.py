from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).parent
output = root / "Smart-Candidate-Screening-System-Report.docx"
document = Document()
section = document.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.color.rgb = RGBColor(23, 59, 99)


def heading(text, level=1):
    paragraph = document.add_heading(text, level=level)
    if level == 1:
        paragraph.runs[0].font.color.rgb = RGBColor(23, 59, 99)
    return paragraph


def paragraph(text=""):
    return document.add_paragraph(text)


def bullets(items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def numbered(items):
    for item in items:
        document.add_paragraph(item, style="List Number")


document.add_heading("Smart Candidate Screening System", 0)
subtitle = document.add_paragraph("Project Report")
subtitle.runs[0].font.size = Pt(14)
subtitle.runs[0].font.color.rgb = RGBColor(83, 101, 121)
paragraph("Purpose: Transparent, consistent, and explainable student screening for placement teams.")
paragraph("Technology: Python, Pandas, Streamlit, and CSV data.")
paragraph("Live application: https://gamepatt78--smart-candidate-screening-system-app-1fibjl.streamlit.app/")
paragraph("GitHub repository: https://github.com/gamepatt78/-smart-candidate-screening-system")
paragraph("Local preview: http://127.0.0.1:8501")

heading("1. Executive Summary")
paragraph("The Smart Candidate Screening System is a transparent web application for helping placement coordinators review student profiles quickly and consistently. It replaces slow, manual spreadsheet screening with a shared rule-based scoring process. Each candidate is evaluated using the same written criteria, and the application explains the reason for every category assigned.")
paragraph("The system supports human decision-making rather than replacing it. Coordinators can filter candidates by identity, branch, CGPA, skills, and readiness category before preparing a recruitment shortlist.")

heading("2. Problem Statement")
paragraph("Manual candidate screening becomes unreliable as the number of students increases. Two coordinators reviewing the same dataset may reach different conclusions because the meaning of a strong candidate is not formally defined. Strong students can be overlooked when reviews are rushed or inconsistent, while recruitment drives often require a shortlist before there is enough time to thoroughly review every student.")
paragraph("Need: A fast, consistent, explainable screening tool that applies shared rules and lets coordinators narrow candidates without manually scanning a spreadsheet.")

heading("3. Project Objectives")
numbered([
    "Create a shared definition of candidate readiness.",
    "Apply the same evaluation rules to every student.",
    "Provide a transparent explanation for each result.",
    "Reduce the time required to prepare recruitment shortlists.",
    "Allow coordinators to locate candidates using practical filters.",
    "Present candidate information in a clear dashboard.",
])

heading("4. Proposed Solution")
paragraph("The application reads candidate information from a CSV file, cleans and standardizes the data, calculates a score for each student, and displays the results in an interactive Streamlit dashboard.")
paragraph("It uses deterministic rules rather than a black-box artificial intelligence model. The same candidate data always produces the same score and category, allowing a coordinator to inspect and explain the result in one sentence.")

heading("5. Evaluation Method")
paragraph("Each candidate can receive a maximum of 9 points:")
table = document.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["Evaluation area", "Scoring rule", "Maximum"]
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text
rows = [
    ("CGPA", "2 points for 8.5 or above; 1 point for 7.0 to 8.49", "2"),
    ("Skills", "2 points for 4 or more skills; 1 point for 2 to 3 skills", "2"),
    ("Projects", "2 points for 3 or more projects; 1 point for 1 to 2 projects", "2"),
    ("Internship", "2 points when internship experience is listed", "2"),
    ("Certifications", "1 point when at least one certification is listed", "1"),
    ("Total", "", "9"),
]
for row in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        cell.text = text
bullets(["Strong: 7 to 9 points", "Average: 4 to 6 points", "Needs Improvement: 0 to 3 points"])

heading("6. Main Features")
heading("Candidate dashboard", 2)
paragraph("Displays the candidate list, total students shown, and counts for each readiness category.")
heading("Search", 2)
paragraph("Searches by candidate name, ID, or branch. Numeric ID searches use exact matching, so searching for ID 2 does not return ID 12.")
heading("CGPA range filtering", 2)
paragraph("The range slider supports minimum and maximum values, such as 3.40 to 6.60. The reset button restores the visible range to 0.0 to 10.0.")
heading("Skills and category filtering", 2)
paragraph("Coordinators can find candidates with skills such as React and combine this with CGPA and category filters.")
heading("Explainable student details", 2)
paragraph("Each candidate has expandable details containing contact information, branch, CGPA, skills, projects, internship experience, certifications, and a scoring explanation.")

heading("7. Technical Implementation")
bullets(["Python: application logic", "Streamlit: interactive web interface", "Pandas: CSV loading, cleaning, filtering, and calculations", "CSV: candidate dataset storage"])
paragraph("The system cleans missing values, converts CGPA values into numbers, separates list fields such as skills and projects, calculates scores, assigns categories, and applies the selected filters.")

heading("8. Benefits")
bullets(["Faster shortlist preparation", "More consistent evaluation across coordinators", "Clear and defensible placement decisions", "Less manual spreadsheet scanning", "Easy combination of academic and technical filters", "Better visibility into each student profile"])

heading("9. Limitations")
paragraph("The system is a screening aid, not a complete recruitment decision-maker. Its results depend on accurate and complete CSV data. The simple scoring rules may not capture communication ability, interview performance, leadership, or role-specific requirements. Authentication, direct CSV upload, shortlist export, recruiter notes, and role-specific scoring are future enhancements.")

heading("10. Future Enhancements")
bullets(["Upload and validate new candidate files", "Export filtered shortlists to CSV or PDF", "Role-specific rules for different recruitment tracks", "Recruiter notes and shortlist status tracking", "Authentication and coordinator-level access control", "Audit history for scoring rule changes"])

heading("11. Conclusion")
paragraph("The Smart Candidate Screening System addresses inconsistency, delay, and lack of explainability in manual candidate review. By applying written scoring rules to every student and providing practical filters, it helps placement coordinators prepare fair shortlists more efficiently.")
paragraph("The system does not make mysterious decisions. It provides a transparent starting point that coordinators can inspect, explain, and use alongside professional judgment.")

heading("12. Screenshots and Product Views")
screenshots = [
    ("dashboard-overview.png", "Dashboard overview with summary metrics and candidate filters."),
    ("student-details.png", "Expandable student details and scoring explanation."),
    ("filter-results.png", "Filtered candidate results."),
    ("scoring-explanation.png", "Transparent scoring rules shown inside the application."),
    ("candidate-list.png", "Candidate list with CGPA, skills, projects, and category."),
]
for filename, caption in screenshots:
    image_path = root / "screenshots" / filename
    document.add_picture(str(image_path), width=Inches(6.7))
    image_caption = document.add_paragraph(caption)
    image_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_caption.runs[0].italic = True
    image_caption.runs[0].font.color.rgb = RGBColor(83, 101, 121)

document.save(output)
print(output)
